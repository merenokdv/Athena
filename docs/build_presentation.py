#!/usr/bin/env python3
"""
Athena — итоговая презентация по макету школы ИИ (СА 102).
Светлый минималистичный product-стиль (ориентир: dribbble.com/glebich).
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent / "Athena_defense.pptx"

# --- Light product palette (Gleb-like: air, soft gray, ink) ---
BG = RGBColor(0xF6, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x14, 0x16, 0x1A)
MUTED = RGBColor(0x6B, 0x73, 0x80)
SOFT = RGBColor(0x9A, 0xA3, 0xAF)
LINE = RGBColor(0xE6, 0xE9, 0xEE)
ACCENT = RGBColor(0x1F, 0x2A, 0x37)  # soft ink accent, not purple
TEAL = RGBColor(0x3A, 0x8F, 0x84)
W = Inches(13.333)
H = Inches(7.5)
MARGIN = Inches(0.7)


def _run(p, text, *, size, bold=False, color=INK, name="Calibri"):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = name
    return r


def _fill(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _soft_shadow(shape):
    """Лёгкая тень — воздух как в product UI."""
    try:
        spPr = shape._element.spPr
        effect = parse_xml(
            f"""<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
              <a:outerShdw blurRad="50800" dist="25400" dir="2700000" algn="tl" rotWithShape="0">
                <a:prstClr val="black"><a:alpha val="8000"/></a:prstClr>
              </a:outerShdw>
            </a:effectLst>"""
        )
        # remove old effectLst if any
        for child in list(spPr):
            if "effectLst" in child.tag:
                spPr.remove(child)
        spPr.append(effect)
    except Exception:
        pass


def blank(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _fill(bg, BG)
    # send bg to back
    spTree = slide.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return slide


def eyebrow(slide, text: str, top=0.35):
    box = slide.shapes.add_textbox(MARGIN, Inches(top), Inches(11.5), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    _run(p, text.upper(), size=11, bold=True, color=SOFT, name="Calibri")
    return box


def h1(slide, text: str, top=0.75, size=32):
    box = slide.shapes.add_textbox(MARGIN, Inches(top), Inches(11.8), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    _run(p, text, size=size, bold=True, color=INK)
    return box


def sub(slide, text: str, top=1.35, size=14):
    box = slide.shapes.add_textbox(MARGIN, Inches(top), Inches(11.8), Inches(0.45))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _run(p, text, size=size, color=MUTED)
    return box


def page(slide, n: int, total: int = 12):
    box = slide.shapes.add_textbox(Inches(12.0), Inches(7.05), Inches(0.9), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    _run(p, f"{n:02d}", size=12, bold=True, color=SOFT)


def card(slide, x, y, w, h, *, shadow=True):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    _fill(sh, WHITE)
    sh.line.color.rgb = LINE
    sh.line.width = Pt(1)
    try:
        sh.adjustments[0] = 0.1
    except Exception:
        pass
    if shadow:
        _soft_shadow(sh)
    return sh


def card_label(slide, x, y, w, text, *, color=TEAL, size=12):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.32))
    p = box.text_frame.paragraphs[0]
    _run(p, text, size=size, bold=True, color=color)
    return box


def card_body(slide, x, y, w, h, text, *, size=14):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        _run(p, line, size=size, color=INK)
    return box


def add_table(slide, rows, left, top, width, height, col_widths=None):
    nrows = len(rows)
    ncols = len(rows[0])
    table_shape = slide.shapes.add_table(nrows, ncols, Inches(left), Inches(top), Inches(width), Inches(height))
    table = table_shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            _run(p, val, size=12, bold=(r == 0), color=INK if r == 0 else MUTED if False else INK)
            # cell fill
            fill = WHITE if r == 0 else BG
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill if r > 0 else RGBColor(0xEE, 0xF1, 0xF5)
            # simplify: header slightly darker
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xEE, 0xF1, 0xF5)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE
    return table_shape


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    total = 12

    # ========== 1 TITLE ==========
    s = blank(prs)
    # soft top accent line
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(0.55), Inches(1.2), Inches(0.06))
    _fill(line, TEAL)

    box = s.shapes.add_textbox(MARGIN, Inches(0.9), Inches(11.5), Inches(1.4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _run(p, "Итоговый проект по программе", size=14, color=MUTED)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(8)
    _run(
        p2,
        "«Системный аналитик промышленных цифровых решений:\nпроектирование, прототипирование, бизнес-эффекты»",
        size=18,
        bold=True,
        color=INK,
    )

    box = s.shapes.add_textbox(MARGIN, Inches(3.0), Inches(11.5), Inches(1.2))
    p = box.text_frame.paragraphs[0]
    _run(p, "Athena", size=44, bold=True, color=INK)
    p2 = box.text_frame.add_paragraph()
    p2.space_before = Pt(6)
    _run(p2, "Локальный DSS-ассистент по базе знаний · RAG + Qwen + UI", size=18, color=MUTED)

    box = s.shapes.add_textbox(MARGIN, Inches(5.3), Inches(7.5), Inches(1.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _run(p, "Поток СА 102", size=13, bold=True, color=INK)
    p2 = tf.add_paragraph()
    _run(p2, "Даты обучения: 31 августа 2026 – 28 сентября 2026", size=12, color=MUTED)
    p3 = tf.add_paragraph()
    _run(p3, "Федеральный проект «Содействие занятости»", size=12, color=MUTED)

    box = s.shapes.add_textbox(Inches(9.2), Inches(5.3), Inches(3.3), Inches(1.0))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    _run(p, "ФИО полностью", size=13, bold=True, color=INK)
    p2 = box.text_frame.add_paragraph()
    p2.alignment = PP_ALIGN.RIGHT
    _run(p2, "подставьте своё имя", size=11, color=SOFT)
    page(s, 1, total)

    # ========== 2 BUSINESS PROBLEM ==========
    s = blank(prs)
    eyebrow(s, "01  /  Бизнес-задача")
    h1(s, "Поиск по заметкам занимает минуты")
    card(s, 0.7, 1.7, 11.9, 1.15, shadow=False)
    card_body(
        s,
        0.95,
        1.85,
        11.4,
        0.9,
        "Специалист вручную ищет ответы в сотнях .md (Obsidian). Облачный чат не видит личную БЗ и уносит данные наружу.",
        size=15,
    )

    card(s, 0.7, 3.15, 5.7, 2.35)
    card_label(s, 0.95, 3.35, 5.2, "Как сейчас")
    card_body(
        s,
        0.95,
        3.75,
        5.2,
        1.5,
        "Открыть vault → поиск по имени\n→ пролистать файлы → собрать ответ.\n10–15 мин на вопрос, легко пропустить\nнужную заметку.",
        size=13,
    )

    card(s, 6.7, 3.15, 5.9, 2.35)
    card_label(s, 6.95, 3.35, 5.4, "Почему важно")
    card_body(
        s,
        6.95,
        3.75,
        5.4,
        1.5,
        "Потери времени knowledge-worker.\nРиск устаревших / неверных шагов.\nНужны источники и отказ вне БЗ.\nТот же сценарий — для отдела.",
        size=13,
    )

    box = s.shapes.add_textbox(MARGIN, Inches(5.8), Inches(11.8), Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _run(p, "Подтверждение: ", size=12, bold=True, color=INK)
    _run(
        p,
        "vault ~956 файлов; ручной поиск 10–15 мин vs целевые 1–3 мин с ассистентом.",
        size=12,
        color=MUTED,
    )
    page(s, 2, total)

    # ========== 3 CONTEXT ==========
    s = blank(prs)
    eyebrow(s, "02  /  Контекст и рамка")
    h1(s, "Что изучаем и что отдаём на выходе", size=28)
    sub(s, "Коротко: персональный knowledge-assistant как пилот DSS для отдела.", top=1.4)

    blocks = [
        (0.7, 2.0, "Контекст", "Личные / учебные заметки\nи сценарий «как я это\nделал раньше?»"),
        (4.85, 2.0, "Аналитическая цель", "Проверить: локальный RAG\nдаёт точные ответы\nсо источниками"),
        (9.0, 2.0, "Стейкхолдеры", "Заказчик: владелец БЗ\nПользователь: специалист\nПриёмка: на защите"),
        (0.7, 4.45, "Входные данные", "956 .md/.txt → 3003 чанка\nChroma + nomic-embed\nпокрытие 100%"),
        (4.85, 4.45, "Ожидаемый результат", "Прототип Athena + UI\nотчёт, тесты, экономика\nзапуск ./scripts/start.sh"),
        (9.0, 4.45, "Внедрение", "Пилот 2–3 недели\ningest при обновлении БЗ\nопция 9b / 27b"),
    ]
    for x, y, lab, body in blocks:
        card(s, x, y, 3.85, 2.15)
        card_label(s, x + 0.22, y + 0.2, 3.4, lab)
        card_body(s, x + 0.22, y + 0.6, 3.4, 1.35, body, size=13)
    page(s, 3, total)

    # ========== 4 AI SOLUTION ==========
    s = blank(prs)
    eyebrow(s, "02  /  Подход")
    h1(s, "Предлагаемое ИИ-решение: Athena")
    items = [
        (0.7, 1.85, "Исходные данные", "Папка заметок ATHENA_KB\n.md / .txt (+ PDF→md)\nлокальный векторный индекс"),
        (6.9, 1.85, "Работа ИИ", "Hybrid retrieve + Qwen\nответ только по контексту\nthink=false, top-k=8"),
        (0.7, 4.25, "Результат для пользователя", "Краткий ответ в UI\nсписок источников (имена)\nмгновенные приветствия"),
        (6.9, 4.25, "Контроль человека", "Человек читает источники\nотказ, если данных нет\nбез авто-действий вовне"),
    ]
    for x, y, lab, body in items:
        card(s, x, y, 5.7, 2.1)
        card_label(s, x + 0.25, y + 0.22, 5.2, lab)
        card_body(s, x + 0.25, y + 0.65, 5.2, 1.25, body, size=14)
    page(s, 4, total)

    # ========== 5 METRICS ==========
    s = blank(prs)
    eyebrow(s, "03  /  Цель и метрики")
    h1(s, "Цель и метрики успеха", size=28)
    card(s, 0.7, 1.7, 11.9, 0.85, shadow=False)
    card_body(
        s,
        0.95,
        1.85,
        11.4,
        0.6,
        "К 28.09.2026 сократить время поиска ответа по БЗ с 10–15 мин до 1–3 мин при точности ≥4/5 и корректных отказах.",
        size=14,
    )
    add_table(
        s,
        [
            ["Показатель", "Сейчас", "Цель", "Как измеряем"],
            ["Время на ответ, мин", "10–15", "1–3", "Секундомер на 5–7 запросах"],
            ["Качество ответа, 1–5", "~2–3 (v1)", "≥4 у ≥70%", "TEST_LOG, экспертная оценка"],
            ["Покрытие индекса", "100%", "100%", "scripts/check_index.py"],
            ["Ложный отказ / галлюцинация", "бывало", "≈0 на тестах", "Кейс вне БЗ + ревью"],
        ],
        0.7,
        2.8,
        11.9,
        3.2,
        col_widths=[3.5, 2.2, 2.5, 3.7],
    )
    page(s, 5, total)

    # ========== 6 CONCEPT ==========
    s = blank(prs)
    eyebrow(s, "04  /  Концепция")
    h1(s, "Концепция проекта")
    # pipeline strip
    steps = [
        ("Ingest", "chunk + embed"),
        ("Chroma", "векторный индекс"),
        ("Retrieve", "hybrid top-k"),
        ("Qwen", "generate"),
        ("UI", "FastAPI"),
    ]
    for i, (a, b) in enumerate(steps):
        x = 0.7 + i * 2.45
        card(s, x, 1.75, 2.3, 1.55)
        card_label(s, x + 0.18, 1.95, 1.95, f"{i+1}. {a}", size=13)
        card_body(s, x + 0.18, 2.4, 1.95, 0.7, b, size=13)

    card(s, 0.7, 3.6, 5.7, 2.85)
    card_label(s, 0.95, 3.8, 5.2, "Стейкхолдеры")
    card_body(
        s,
        0.95,
        4.25,
        5.2,
        1.9,
        "Заказчик — владелец базы знаний\nПользователь — knowledge-worker\nИТ/ИБ — локальный контур данных\nНа защите — комиссия программы",
        size=14,
    )
    card(s, 6.7, 3.6, 5.9, 2.85)
    card_label(s, 6.95, 3.8, 5.4, "Продукты на выходе")
    card_body(
        s,
        6.95,
        4.25,
        5.4,
        1.9,
        "Рабочий прототип Athena\nОтчёт REPORT.md + тест-лог\nСкрипты ingest / check_index\nПрезентация и живое демо",
        size=14,
    )
    page(s, 6, total)

    # ========== 7 PILOT BOUNDARIES ==========
    s = blank(prs)
    eyebrow(s, "05  /  Границы пилота")
    h1(s, "Границы пилота и готовность данных")
    card(s, 0.7, 1.75, 5.7, 2.4)
    card_label(s, 0.95, 1.95, 5.2, "Входит в пилот")
    card_body(
        s,
        0.95,
        2.4,
        5.2,
        1.5,
        "1 пользователь / свой vault\nВопросы по заметкам (.md/.txt)\nЛокальный UI http://127.0.0.1:7860\nМодели 9b (скорость) / 27b",
        size=14,
    )
    card(s, 6.7, 1.75, 5.9, 2.4)
    card_label(s, 6.95, 1.95, 5.4, "За пределами")
    card_body(
        s,
        6.95,
        2.4,
        5.4,
        1.5,
        "Мультипользователь / SSO\nЮр./мед. заключения «от себя»\nОблачный share заметок\nАвто-изменения в системах",
        size=14,
    )
    add_table(
        s,
        [
            ["Данные", "Готовность и действия"],
            ["956 файлов vault, md/txt", "Доступ есть, индекс собран"],
            ["Качество / шум PDF", "PDF→md, hybrid retrieve, re-ingest"],
        ],
        0.7,
        4.4,
        11.9,
        1.6,
        col_widths=[5.5, 6.4],
    )
    page(s, 7, total)

    # ========== 8 TRIAL / MVP ==========
    s = blank(prs)
    eyebrow(s, "06  /  Пробное действие")
    h1(s, "MVP и цикл улучшения")
    cards3 = [
        (0.7, "Алгоритм", "Ingest → hybrid retrieve\n→ Qwen по контексту\n→ источники в UI"),
        (4.85, "MVP", "Athena UI + Chroma\nqwen3.5:9b/27b\nstart.sh одной командой"),
        (9.0, "Рефлексия", "v1: шум PDF\nv2: чистый md\nv3+: hybrid + ключи"),
    ]
    for x, lab, body in cards3:
        card(s, x, 1.75, 3.85, 2.5)
        card_label(s, x + 0.22, 1.95, 3.4, lab)
        card_body(s, x + 0.22, 2.45, 3.4, 1.5, body, size=14)

    card(s, 0.7, 4.55, 11.9, 1.9, shadow=False)
    card_label(s, 0.95, 4.75, 11.4, "Что получилось / что дальше")
    card_body(
        s,
        0.95,
        5.2,
        11.4,
        1.0,
        "Получилось: 100% покрытие, живой UI, отказ вне БЗ, ускорение ответов на 9b.\n"
        "Дальше: заполнить TEST_LOG на защите, пилот 20 реальных вопросов, регламент re-ingest.",
        size=14,
    )
    page(s, 8, total)

    # ========== 9 TEAM ==========
    s = blank(prs)
    eyebrow(s, "07  /  Команда и роли")
    h1(s, "Команда и ответственность")
    add_table(
        s,
        [
            ["Роль", "Участник", "Ответственность"],
            ["Спонсор / заказчик", "Владелец БЗ (слушатель)", "Цели пилота, решение о продолжении"],
            ["Владелец задачи", "Слушатель СА 102", "Требования, приёмка, отчёт"],
            ["ИИ / данные", "Слушатель (прототип)", "RAG, UI, индекс, промпт"],
            ["Пользователь пилота", "Тот же / коллега", "Тест-запросы и feedback"],
        ],
        0.7,
        1.8,
        11.9,
        3.6,
        col_widths=[3.2, 3.8, 4.9],
    )
    box = s.shapes.add_textbox(MARGIN, Inches(5.8), Inches(11.8), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    _run(p, "Ответственный за пилот: ", size=13, bold=True, color=INK)
    _run(p, "слушатель (подставьте ФИО). Согласования ИТ/ИБ: локальный контур, без облака.", size=13, color=MUTED)
    page(s, 9, total)

    # ========== 10 TIMELINE ==========
    s = blank(prs)
    eyebrow(s, "08  /  Сроки и ресурсы")
    h1(s, "Сроки и ресурсы")
    phases = [
        ("Неделя 1", "Железо, Ollama, GPU\nкаркас RAG + ingest"),
        ("Неделя 2", "UI Athena, hybrid\nдоработки качества"),
        ("Неделя 3", "Тесты, экономика\nотчёт и защита"),
        ("Пилот+", "20 вопросов, feedback\nрегламент re-ingest"),
    ]
    for i, (lab, body) in enumerate(phases):
        x = 0.7 + i * 3.1
        card(s, x, 1.8, 2.95, 2.8)
        card_label(s, x + 0.2, 2.05, 2.55, lab)
        card_body(s, x + 0.2, 2.6, 2.55, 1.7, body, size=14)

    card(s, 0.7, 5.0, 11.9, 1.5, shadow=False)
    card_label(s, 0.95, 5.2, 11.4, "Ресурсы")
    card_body(
        s,
        0.95,
        5.6,
        11.4,
        0.7,
        "RTX 3090 24 ГБ · ~62 ГБ RAM · Ollama · лицензии 0 ₽ · разработка прототипа 12–20 ч · поддержка 2–4 ч/мес",
        size=14,
    )
    page(s, 10, total)

    # ========== 11 RISKS ==========
    s = blank(prs)
    eyebrow(s, "09  /  Риски и решение")
    h1(s, "Риски и критерии решения", size=28)
    add_table(
        s,
        [
            ["Ключевой риск", "Как снижаем", "Ответственный"],
            ["Неполная / устаревшая БЗ", "ingest + check_index после правок", "Владелец БЗ"],
            ["Шумный retrieval", "hybrid, PDF→md, diversity top-k", "ИИ/данные"],
            ["Галлюцинации", "prompt «только контекст» + источники", "ИИ/данные"],
            ["Медленные ответы", "GPU, 9b, think=false, unload VRAM", "ИИ/данные"],
        ],
        0.7,
        1.7,
        11.9,
        3.0,
        col_widths=[3.8, 5.0, 3.1],
    )
    decisions = [
        (0.7, "Масштабировать", "≥70% ответов ≥4\nотказы корректны\nпользователь экономит время"),
        (4.85, "Доработать", "Качество 3–4\nузкие темы / имена файлов\nдожать retrieve и корпус"),
        (9.0, "Остановить", "Крит. галлюцинации\nнет эффекта по времени\nданные нельзя держать локально"),
    ]
    for x, lab, body in decisions:
        card(s, x, 5.0, 3.85, 1.7)
        card_label(s, x + 0.2, 5.15, 3.4, lab, size=12)
        card_body(s, x + 0.2, 5.5, 3.4, 1.0, body, size=12)
    page(s, 11, total)

    # ========== 12 THANKS ==========
    s = blank(prs)
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(2.6), Inches(1.0), Inches(0.06))
    _fill(line, TEAL)
    box = s.shapes.add_textbox(MARGIN, Inches(2.9), Inches(11.5), Inches(1.5))
    p = box.text_frame.paragraphs[0]
    _run(p, "Спасибо за внимание", size=40, bold=True, color=INK)
    p2 = box.text_frame.add_paragraph()
    p2.space_before = Pt(12)
    _run(p2, "Athena · демо http://127.0.0.1:7860 · вопросы?", size=16, color=MUTED)
    box = s.shapes.add_textbox(MARGIN, Inches(5.5), Inches(11.5), Inches(0.8))
    p = box.text_frame.paragraphs[0]
    _run(p, "Поток СА 102  ·  31.08–28.09.2026", size=13, color=SOFT)
    page(s, 12, total)

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
